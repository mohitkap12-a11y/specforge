"""Spec -> Markdown converter: PDF and DOCX structure, escaping, the API endpoint and the CLI."""
import io
import re
import shutil
import zipfile

import docx
import pytest
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from fastapi.testclient import TestClient

from specforge.convert import ConversionError, convert_file, output_name, sniff
from specforge.convert import md


# ----------------------------------------------------------------------------- fixtures

@pytest.fixture(scope="module")
def rich_docx(tmp_path_factory):
    d = docx.Document()
    d.core_properties.title = "Widget Spec"
    d.add_heading("Widget Interface", 0)
    d.add_heading("Overview", 1)
    p = d.add_paragraph("A widget ")
    p.add_run("shall").bold = True
    p.add_run(" answer ")
    p.add_run("GET_ID").font.name = "Courier New"
    p.add_run(" within 10 ms; see ")
    rid = d.part.relate_to("https://example.com/a b", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "the site"
    r.append(t)
    h.append(r)
    p._p.append(h)
    p.add_run(". Value 2")
    p.add_run("n").font.superscript = True
    d.add_heading("Details", 2)
    d.add_paragraph("first bullet", style="List Bullet")
    d.add_paragraph("nested bullet", style="List Bullet 2")
    d.add_paragraph("step one", style="List Number")
    d.add_paragraph("step two", style="List Number")
    d.add_paragraph("snake_case_name and * star")
    d.add_paragraph("# not a heading")
    t = d.add_table(rows=3, cols=3)
    for c, v in zip(t.rows[0].cells, ("Bits", "Name", "Desc")):
        c.text = v
    t.cell(1, 0).text, t.cell(1, 1).text, t.cell(1, 2).text = "7:0", "A|B", "line1"
    t.cell(1, 2).add_paragraph("line2")
    t.cell(2, 0).merge(t.cell(2, 1)).text = "merged"
    t.cell(2, 2).text = "x"
    path = tmp_path_factory.mktemp("conv") / "rich.docx"
    d.save(path)
    return path


@pytest.fixture(scope="module")
def rich_pdf(tmp_path_factory):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Table, TableStyle

    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Normal"], fontSize=20, leading=24, fontName="Helvetica-Bold")
    h2 = ParagraphStyle("h2", parent=ss["Normal"], fontSize=15, leading=18, fontName="Helvetica-Bold")
    n = ParagraphStyle("n", parent=ss["Normal"], fontSize=10, leading=13)
    code = ParagraphStyle("c", parent=ss["Normal"], fontName="Courier", fontSize=9, leading=11)

    def furniture(c, doc):
        c.setFont("Helvetica", 8)
        c.drawString(50, 820, "ACME Confidential Spec")
        c.drawString(300, 20, f"Page {doc.page}")

    rows = [["Reg", "Desc"]] + [[f"R{i}", f"Register number {i}"] for i in range(70)]
    story = [Paragraph("Widget Spec", h1), Paragraph("1 Scope", h2),
             Paragraph("The controller shall be conservative in what it sends and liberal in what it accepts "
                       "from the host, including <b>bold text</b> and <i>italic</i> words. " * 3, n),
             Paragraph("• first item", n), Paragraph("• second item", n),
             Paragraph("1. numbered one", n), Paragraph("2. numbered two", n),
             Preformatted("int x = 1;\n  if (x) y();", code),
             Paragraph("2 Registers", h2),
             Table(rows, repeatRows=1, style=TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)])),
             PageBreak(), Paragraph("3 Tail", h2), Paragraph("The end.", n)]
    path = tmp_path_factory.mktemp("convpdf") / "rich.pdf"
    SimpleDocTemplate(str(path), pagesize=A4).build(story, onFirstPage=furniture, onLaterPages=furniture)
    return path


# ----------------------------------------------------------------------------- markdown helpers

def test_escape_keeps_snake_case_and_escapes_markup():
    assert md.escape_text("snake_case and _lead and trail_ *x* <b>") == "snake_case and \\_lead and trail\\_ \\*x\\* \\<b>"


def test_inline_moves_whitespace_outside_emphasis_and_joins_phrases():
    segs = [md.Seg("a ", bold=True), md.Seg("b ", bold=True), md.Seg("c")]
    assert md.inline(segs) == "**a b** c"
    assert md.inline([md.Seg("x", bold=True), md.Seg(" "), md.Seg("y", bold=True)]) == "**x y**"


def test_paragraph_that_looks_like_markdown_is_escaped():
    out = md.render([md.Para([md.Seg("# not a heading")]), md.Para([md.Seg("1. not a list")]), md.Para([md.Seg("- dash")])])
    assert out.split("\n\n") == ["\\# not a heading", "1\\. not a list", "\\- dash\n"]


def test_table_escapes_pipes_and_pads_rows():
    assert md.table_md([["a", "b|c"], ["1"]]) == "| a | b\\|c |\n| --- | --- |\n| 1 |  |"


def test_lists_stay_tight_and_nested():
    out = md.render([md.Item(0, [md.Seg("a")]), md.Item(1, [md.Seg("b")]), md.Item(0, [md.Seg("c")], True, "1"), md.Para([md.Seg("after")])])
    assert out == "- a\n  - b\n\n1. c\n\nafter\n"


# ----------------------------------------------------------------------------- DOCX

def test_docx_structure(rich_docx):
    r = convert_file(rich_docx, front_matter=False)
    m = r.markdown
    assert "# Widget Interface" in m and "# Overview" in m and "## Details" in m
    assert "**shall**" in m and "`GET_ID`" in m
    assert "[the site](https://example.com/a%20b)" in m
    assert "2<sup>n</sup>" in m
    assert "- first bullet\n  - nested bullet" in m
    assert "1. step one\n2. step two" in m
    assert "snake_case_name and \\* star" in m
    assert "\\# not a heading" in m
    assert "| 7:0 | A\\|B | line1<br>line2 |" in m
    assert "| merged |  | x |" in m
    assert r.stats["tables"] == 1 and r.stats["headings"] == 3 and r.title == "Widget Spec"


def test_docx_front_matter(rich_docx):
    m = convert_file(rich_docx, filename="Widget v1.docx").markdown
    assert m.startswith('---\ntitle: "Widget Spec"\nsource: "Widget v1.docx"\nformat: docx\n---\n')


def test_docx_numbered_headings_and_footnotes(tmp_path):
    """Word stores heading numbers in numbering.xml, not in the text; footnotes live in their own part."""
    d = docx.Document()
    d.add_heading("Scope", 1)
    d.add_heading("Terms", 2)
    d.add_heading("Other", 1)
    d.add_heading("2-bit flags", 1)
    path = tmp_path / "num.docx"
    d.save(path)
    # give Heading 1 / Heading 2 a multilevel numbering definition
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    numbering = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                 b'<w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading1"/><w:lvlText w:val="%1"/></w:lvl>'
                 b'<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading2"/><w:lvlText w:val="%1.%2"/></w:lvl></w:abstractNum>'
                 b'<w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>')
    files["word/numbering.xml"] = numbering
    styles = files["word/styles.xml"].decode()
    for sid in ("Heading1", "Heading2"):
        marker = f'w:styleId="{sid}"'
        i = styles.index(marker)
        j = styles.index("<w:pPr>", i)
        styles = styles[:j] + '<w:pPr><w:numPr><w:numId w:val="1"/></w:numPr>' + styles[j + len("<w:pPr>"):]
    files["word/styles.xml"] = styles.encode()
    with zipfile.ZipFile(path, "w") as z:
        for n, b in files.items():
            z.writestr(n, b)
    m = convert_file(path, front_matter=False).markdown
    assert "# 1 Scope" in m and "## 1.1 Terms" in m and "# 2 Other" in m
    assert "# 3 2-bit flags" in m   # a label is only skipped when the text really starts with it


def test_docx_corrupt_and_wrong_type(tmp_path):
    bad = tmp_path / "bad.docx"
    with zipfile.ZipFile(bad, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(ConversionError):
        convert_file(bad)
    txt = tmp_path / "a.pdf"
    txt.write_text("plain text pretending to be a pdf")
    with pytest.raises(ConversionError, match="Unsupported"):
        convert_file(txt)
    with pytest.raises(ConversionError, match="Cannot read"):
        sniff(tmp_path / "missing.pdf")


# ----------------------------------------------------------------------------- PDF

def test_pdf_structure(rich_pdf):
    r = convert_file(rich_pdf, front_matter=False)
    m = r.markdown
    assert m.startswith("# Widget Spec\n\n## 1 Scope\n")
    assert "## 2 Registers" in m and "## 3 Tail" in m
    assert "**bold text**" in m and "*italic*" in m
    assert "- first item\n- second item" in m
    assert "1. numbered one\n2. numbered two" in m
    assert "```\nint x = 1;\n  if (x) y();\n```" in m or "```\nint x = 1;\nif (x) y();\n```" in m
    assert "ACME Confidential" not in m and "Page 1" not in m          # running header / footer removed
    assert r.stats["pages"] >= 2 and r.stats["tables"] == 1
    # the table that spans pages is one table with a single header and all 70 rows
    assert m.count("| Reg | Desc |") == 1 and "| R69 | Register number 69 |" in m and len(re.findall(r"^\| R\d+ \|", m, re.M)) == 70


def test_pdf_hyphenated_line_break_is_healed(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "A controller shall support the inter-", fontsize=11)
    page.insert_text((72, 114), "operability requirements of the host.", fontsize=11)
    path = tmp_path / "hy.pdf"
    doc.save(path)
    assert "interoperability requirements" in convert_file(path, front_matter=False).markdown


def test_pdf_bookmarks_set_heading_levels(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    for t in ("Intro", "Details"):
        page = doc.new_page()
        page.insert_text((72, 100), t, fontsize=11)
        page.insert_text((72, 130), "Body text of the section.", fontsize=11)
    doc.set_toc([[1, "Intro", 1], [2, "Details", 2]])
    path = tmp_path / "bm.pdf"
    doc.save(path)
    m = convert_file(path, front_matter=False).markdown
    assert "# Intro\n" in m and "## Details\n" in m


def test_pdf_without_text_layer_is_rejected(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    doc.new_page()
    path = tmp_path / "blank.pdf"
    doc.save(path)
    with pytest.raises(ConversionError, match="no text layer"):
        convert_file(path)


def test_demo_specs_convert_with_all_tables(demo_dir):
    for name in ("nvme-base-2.4.pdf", "nvme-base-2.4.docx"):
        r = convert_file(demo_dir / name, front_matter=False)
        assert "| 01h | M | Create I/O Submission Queue |" in r.markdown
        assert "## 5.1 Admin Command Overview" in r.markdown
        assert r.stats["tables"] == 11


# ----------------------------------------------------------------------------- .doc, API, CLI

@pytest.mark.skipif(not (shutil.which("soffice") or shutil.which("libreoffice")), reason="LibreOffice not installed")
def test_legacy_doc_via_libreoffice(rich_docx, tmp_path):
    import subprocess

    subprocess.run(["soffice", "--headless", "--convert-to", "doc:MS Word 97", "--outdir", str(tmp_path), str(rich_docx)], check=True, capture_output=True, timeout=180)
    doc = tmp_path / "rich.doc"
    assert sniff(doc) == "doc"
    assert "## Details" in convert_file(doc, front_matter=False).markdown


def test_doc_without_libreoffice_gives_a_clear_error(tmp_path, monkeypatch):
    monkeypatch.setattr(shutil, "which", lambda _: None)
    f = tmp_path / "old.doc"
    f.write_bytes(b"\xd0\xcf\x11\xe0" + b"\0" * 64)
    with pytest.raises(ConversionError, match="LibreOffice"):
        convert_file(f)


def test_output_name_is_download_safe():
    assert output_name("NVM Express Base 2.4.pdf") == "NVM Express Base 2.4.md"
    assert output_name('we"ird/\\na.docx') == "na.md" or output_name('we"ird/\\na.docx').endswith(".md")
    assert output_name("") == "document.md"


@pytest.fixture()
def client(ws):
    from specforge.api.app import create_app

    with TestClient(create_app()) as c:
        yield c


def test_api_convert_returns_markdown_and_stores_nothing(client, rich_docx, rich_pdf):
    with open(rich_docx, "rb") as f:
        r = client.post("/api/convert/markdown", files={"file": ("Widget Spec.docx", f)})
    assert r.status_code == 200
    j = r.json()
    assert j["filename"] == "Widget Spec.md" and j["format"] == "docx" and "## Details" in j["markdown"] and j["markdown"].startswith("---")
    with open(rich_pdf, "rb") as f:
        r = client.post("/api/convert/markdown?front_matter=false", files={"file": ("w.pdf", f)})
    assert r.json()["markdown"].startswith("# Widget Spec")
    assert client.get("/api/specs").json() == []


def test_api_convert_download_and_errors(client, rich_docx):
    with open(rich_docx, "rb") as f:
        r = client.post("/api/convert/markdown?download=true", files={"file": ("Spéc.docx", f)})
    assert r.headers["content-type"].startswith("text/markdown")
    assert "filename*=UTF-8''Sp%C3%A9c.md" in r.headers["content-disposition"]
    assert r.text.startswith("---")
    assert client.post("/api/convert/markdown", files={"file": ("a.txt", io.BytesIO(b"hi"))}).status_code == 400
    bad = client.post("/api/convert/markdown", files={"file": ("a.pdf", io.BytesIO(b"not a pdf"))})
    assert bad.status_code == 422 and "Unsupported" in bad.json()["detail"]


def test_cli_to_md(rich_docx, rich_pdf, tmp_path):
    from specforge.cli import main

    main(["to-md", str(rich_docx), str(rich_pdf), "--out", str(tmp_path / "out")])
    assert {p.name for p in (tmp_path / "out").iterdir()} == {"rich.md", "rich.pdf.md"}
    with pytest.raises(SystemExit) as e:
        main(["to-md", str(tmp_path / "nope.pdf")])
    assert e.value.code == 1


# ----------------------------------------------------------------------------- review regressions

def test_code_keeps_literal_spacing_blank_lines_and_backticks():
    assert md.inline([md.Seg("a  b", code=True)]) == "`a  b`"
    assert md.inline([md.Seg("a`b``c", code=True)]) == "```a`b``c```"
    out = md.render([md.Code(["x = 1", "", "", "y = 2"]), md.Code(["````", "z"])])
    assert "x = 1\n\n\ny = 2" in out and "`````\n````\nz\n`````" in out
    assert md.inline([md.Seg("two  spaces in prose")]) == "two spaces in prose"


def test_pdf_distinct_tables_with_equal_columns_are_not_merged(tmp_path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import PageBreak, SimpleDocTemplate, Table, TableStyle

    grid = TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)])
    path = tmp_path / "two.pdf"
    SimpleDocTemplate(str(path), pagesize=A4).build([
        Table([["Reg", "Desc"], ["R0", "zero"]], style=grid), PageBreak(),
        Table([["Opcode", "Name"], ["01h", "Write"]], style=grid)])
    m = convert_file(path, front_matter=False).markdown
    assert m.count("| --- | --- |") == 2 and "| Opcode | Name |" in m and "| R0 | zero |" in m


def test_pdf_list_nesting_is_consistent_across_pages(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    p1 = doc.new_page()
    p1.insert_text((72, 100), "Introduction text for the list below.", fontsize=11)
    p1.insert_text((72, 130), "• parent item", fontsize=11)
    p2 = doc.new_page()
    p2.insert_text((90, 100), "• child item", fontsize=11)
    p2.insert_text((72, 130), "• other parent", fontsize=11)
    path = tmp_path / "nest.pdf"
    doc.save(path)
    m = convert_file(path, front_matter=False).markdown
    lines = [ln for ln in m.splitlines() if "item" in ln or "parent" in ln]
    assert lines[0].startswith("- parent") and lines[1].startswith("  - child") and lines[2].startswith("- other")


def test_pdf_non_latin_bookmarks_do_not_turn_body_lines_into_headings(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Plain body line one.", fontsize=11)
    doc.set_toc([[1, "概要", 1]])
    path = tmp_path / "jp.pdf"
    doc.save(path)
    assert convert_file(path, front_matter=False).markdown.strip() == "Plain body line one."


def test_cli_never_overwrites_an_input_or_an_earlier_output(tmp_path):
    from specforge.cli import _free_target

    src = tmp_path / "spec.md"
    written = set()
    t1 = _free_target(tmp_path / "spec.md", "spec.md", written, {src.resolve()})
    assert t1 != src
    written.add(t1)
    t2 = _free_target(tmp_path / "spec.md", "spec.md", written, {src.resolve()})
    t3 = _free_target(tmp_path / "spec.md", "spec.md", written | {t2}, {src.resolve()})
    assert len({src, t1, t2, t3}) == 4


def test_docx_huge_roman_start_is_bounded():
    from specforge.convert.word import _format_num

    assert _format_num(10**15, "upperRoman") == str(10**15)
    assert _format_num(4, "lowerRoman") == "iv"


def test_docx_tabs_survive_in_code_paragraphs(tmp_path):
    d = docx.Document()
    p = d.add_paragraph()
    r = p.add_run("if (x)\n\ty();")
    r.font.name = "Courier New"
    path = tmp_path / "tab.docx"
    d.save(path)
    assert "if (x)\n\ty();" in convert_file(path, front_matter=False).markdown


def test_docx_huge_grid_span_is_clamped(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "a", "b", "c", "d"
    tc = t.cell(0, 0)._tc
    pr = tc.get_or_add_tcPr()
    gs = OxmlElement("w:gridSpan")
    gs.set(qn("w:val"), "2147483647")
    pr.append(gs)
    path = tmp_path / "span.docx"
    d.save(path)
    assert "| c | d |" in convert_file(path, front_matter=False).markdown


def test_paragraph_lines_after_a_manual_break_are_escaped():
    out = md.render([md.Para([md.Seg("first\n# second\n- third")])])
    assert out == "first  \n\\# second  \n\\- third\n"


def test_pdf_table_cells_are_literal(tmp_path):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle

    path = tmp_path / "lit.pdf"
    SimpleDocTemplate(str(path), pagesize=A4).build([Table([["Field", "Note"], ["a*b*", "<NSID> `x`"]], style=TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black)]))])
    m = convert_file(path, front_matter=False).markdown
    assert "a\\*b\\*" in m and "\\<NSID> \\`x\\`" in m


def test_docx_numbered_cell_paragraphs_advance_counters_and_override_levels_apply(tmp_path):
    d = docx.Document()
    d.add_paragraph("before", style="List Number")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "k"
    t.cell(0, 1).paragraphs[0].text = "in cell"
    t.cell(0, 1).paragraphs[0].style = d.styles["List Number"]
    d.add_paragraph("after", style="List Number")
    path = tmp_path / "cell.docx"
    d.save(path)
    m = convert_file(path, front_matter=False).markdown
    assert "| k | 2. in cell |" in m and "3. after" in m


def test_brackets_and_entities_stay_literal_but_footnote_refs_work():
    assert md.inline([md.Seg("[7:0]: Reserved &copy; 5 & 6")]) == "\\[7:0\\]: Reserved &amp;copy; 5 & 6"
    assert md.inline([md.Seg("text"), md.Seg("[^1]", raw=True)]) == "text[^1]"
    assert md.render([md.Para([md.Seg("[7:0]: Reserved")])]) == "\\[7:0\\]: Reserved\n"


def test_docx_footnote_reference_is_emitted_with_definition(tmp_path):
    from specforge.convert import word

    d = docx.Document()
    d.add_paragraph("Body.")
    path = tmp_path / "fn.docx"
    d.save(path)
    ctx = word._Ctx(docx.Document(path))
    ctx.footnotes["2"] = "Note text"
    r = OxmlElement("w:r")
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), "2")
    r.append(ref)
    assert md.inline(ctx.run(r, None, None, False)) == "[^2]"


def test_nested_items_indent_past_the_parent_marker():
    out = md.render([md.Item(0, [md.Seg("p")], True, "100"), md.Item(1, [md.Seg("c")]), md.Item(2, [md.Seg("g")])])
    assert out == "100. p\n     - c\n       - g\n"
    # an item whose parent is missing nests under the deepest open level instead of detaching
    assert md.render([md.Item(3, [md.Seg("x")])]) == "- x\n"


def test_heading_with_trailing_hashes_keeps_them():
    assert md.render([md.Heading(2, "Language C #")]) == "## Language C \\#\n"


def test_docx_extreme_list_depth_is_clamped(tmp_path):
    d = docx.Document()
    p = d.add_paragraph("deep", style="List Bullet")
    ilvl = p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_ilvl()
    ilvl.set(qn("w:val"), "2147483647")
    path = tmp_path / "deep.docx"
    d.save(path)
    assert len(convert_file(path, front_matter=False).markdown) < 200


def test_docx_contents_heading_is_kept_unless_it_held_a_toc(tmp_path):
    d = docx.Document()
    d.add_heading("Contents", 1)
    d.add_heading("Hardware", 2)
    path = tmp_path / "c.docx"
    d.save(path)
    assert "# Contents" in convert_file(path, front_matter=False).markdown


def test_pdf_paragraph_after_list_is_not_swallowed(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "\u2022 only item", fontsize=11)
    page.insert_text((72, 114), "A new paragraph at the margin.", fontsize=11)
    path = tmp_path / "li.pdf"
    doc.save(path)
    assert "- only item\n\nA new paragraph at the margin." in convert_file(path, front_matter=False).markdown


def test_docx_monospace_list_items_stay_list_items(tmp_path):
    d = docx.Document()
    for t in ("first_cmd", "second_cmd"):
        d.add_paragraph(style="List Number").add_run(t).font.name = "Courier New"
    path = tmp_path / "ml.docx"
    d.save(path)
    m = convert_file(path, front_matter=False).markdown
    assert "1. `first_cmd`\n2. `second_cmd`" in m


def test_docx_hyperlink_anchor_is_kept(tmp_path):
    d = docx.Document()
    p = d.add_paragraph("see ")
    rid = d.part.relate_to("https://example.com/doc", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), rid)
    h.set(qn("w:anchor"), "sec2")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "here"
    r.append(t)
    h.append(r)
    p._p.append(h)
    path = tmp_path / "anchor.docx"
    d.save(path)
    assert "[here](https://example.com/doc#sec2)" in convert_file(path, front_matter=False).markdown


def test_superscript_segments_render_as_sup():
    assert md.inline([md.Seg("2"), md.Seg("n", sup=True)]) == "2<sup>n</sup>"


def test_pdf_adjacent_bookmarked_headings_stay_separate(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Alpha", fontsize=11)
    page.insert_text((72, 112), "Beta", fontsize=11)
    doc.set_toc([[1, "Alpha", 1], [1, "Beta", 1]])
    path = tmp_path / "two.pdf"
    doc.save(path)
    m = convert_file(path, front_matter=False).markdown
    assert "# Alpha\n\n# Beta" in m


def test_docx_grid_before_keeps_values_under_their_columns(tmp_path):
    d = docx.Document()
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "A", "B", "x", "y"
    t.cell(1, 0)._tc.getparent().remove(t.cell(1, 0)._tc)   # row 2 now starts one column in
    tr = t.rows[1]._tr
    pr = tr.get_or_add_trPr()
    gb = OxmlElement("w:gridBefore")
    gb.set(qn("w:val"), "1")
    pr.append(gb)
    path = tmp_path / "gb.docx"
    d.save(path)
    assert "|  | y |" in convert_file(path, front_matter=False).markdown


def test_pdf_code_indentation_from_leading_spaces_is_kept(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "def f():", fontname="cour", fontsize=10)
    page.insert_text((72, 112), "    return 1", fontname="cour", fontsize=10)
    path = tmp_path / "code.pdf"
    doc.save(path)
    assert "def f():\n    return 1" in convert_file(path, front_matter=False).markdown


def test_list_and_quote_text_that_looks_like_markdown_stays_literal():
    out = md.render([md.Item(0, [md.Seg("# Foo")]), md.Item(0, [md.Seg("- Bar")]), md.Quote([md.Seg("# Baz")])])
    assert out == "- \\# Foo\n- \\- Bar\n\n> \\# Baz\n"


def test_tilde_fence_at_line_start_is_escaped():
    assert md.render([md.Para([md.Seg("~~~ not a fence")])]) == "\\~~~ not a fence\n"


def test_link_destination_parentheses_are_encoded():
    assert md.inline([md.Seg("x", link="https://e.com/a(b)c")]) == "[x](https://e.com/a%28b%29c)"


def test_docx_local_ilvl_only_keeps_the_style_list(tmp_path):
    d = docx.Document()
    d.add_paragraph("one", style="List Number")
    p = d.add_paragraph("two", style="List Number")
    p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_ilvl().set(qn("w:val"), "0")
    d.add_paragraph("three", style="List Number")
    path = tmp_path / "ilvl.docx"
    d.save(path)
    assert "1. one\n2. two\n3. three" in convert_file(path, front_matter=False).markdown


def _ctx_with_numbering(tmp_path, abstract_levels: str):
    d = docx.Document()
    d.add_paragraph("x")
    path = tmp_path / "n.docx"
    d.save(path)
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    files["word/numbering.xml"] = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                                   f'<w:abstractNum w:abstractNumId="0">{abstract_levels}</w:abstractNum><w:num w:numId="1"><w:abstractNumId w:val="0"/></w:num></w:numbering>').encode()
    with zipfile.ZipFile(path, "w") as z:
        for n, b in files.items():
            z.writestr(n, b)
    from specforge.convert import word

    return word._Ctx(docx.Document(path))


def test_docx_lvl_restart_zero_keeps_child_counter(tmp_path):
    lv0 = '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
    lv1 = '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlRestart w:val="0"/><w:lvlText w:val="%2)"/></w:lvl>'
    ctx = _ctx_with_numbering(tmp_path, lv0 + lv1)
    labels = [ctx.label("1", i)[0] for i in (0, 1, 0, 1)]
    assert labels == ["1.", "1)", "2.", "2)"]   # the child level never restarts
    lv1 = lv1.replace('<w:lvlRestart w:val="0"/>', "")
    ctx = _ctx_with_numbering(tmp_path, lv0 + lv1)
    assert [ctx.label("1", i)[0] for i in (0, 1, 0, 1)] == ["1.", "1)", "2.", "1)"]


def test_docx_style_linked_level_wins_over_style_ilvl(tmp_path):
    lv0 = '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading1"/><w:lvlText w:val="%1"/></w:lvl>'
    lv1 = '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:pStyle w:val="Heading2"/><w:lvlText w:val="%1.%2"/></w:lvl>'
    ctx = _ctx_with_numbering(tmp_path, lv0 + lv1)
    assert ctx.style_level[("1", "Heading2")] == 1


def test_docx_footnotes_do_not_expand_external_entities(tmp_path):
    secret = tmp_path / "secret.txt"
    secret.write_text("TOP-SECRET")
    d = docx.Document()
    d.add_paragraph("Body.")
    path = tmp_path / "xxe.docx"
    d.save(path)
    with zipfile.ZipFile(path) as z:
        files = {n: z.read(n) for n in z.namelist()}
    files["word/footnotes.xml"] = (
        f'<?xml version="1.0"?><!DOCTYPE r [<!ENTITY x SYSTEM "file://{secret}">]>'
        '<w:footnotes xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:footnote w:id="2"><w:p><w:r><w:t>&x;</w:t></w:r></w:p></w:footnote></w:footnotes>').encode()
    files["[Content_Types].xml"] = files["[Content_Types].xml"].replace(b"</Types>", b'<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
    files["word/_rels/document.xml.rels"] = files["word/_rels/document.xml.rels"].replace(
        b"</Relationships>", b'<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
    with zipfile.ZipFile(path, "w") as z:
        for n, b in files.items():
            z.writestr(n, b)
    from specforge.convert import word

    ctx = word._Ctx(docx.Document(path))
    assert "2" in ctx.footnotes and "TOP-SECRET" not in ctx.footnotes["2"]
    assert "TOP-SECRET" not in convert_file(path, front_matter=False).markdown


def test_docx_blank_paragraph_inside_code_keeps_one_block(tmp_path):
    d = docx.Document()
    for t in ("a = 1", "", "b = 2"):
        p = d.add_paragraph()
        if t:
            p.add_run(t).font.name = "Courier New"
    path = tmp_path / "cb.docx"
    d.save(path)
    assert "```\na = 1\n\nb = 2\n```" in convert_file(path, front_matter=False).markdown


def test_docx_tracked_move_destination_text_is_kept(tmp_path):
    d = docx.Document()
    p = d.add_paragraph("keep ")
    mv = OxmlElement("w:moveTo")
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = "moved"
    r.append(t)
    mv.append(r)
    p._p.append(mv)
    path = tmp_path / "mv.docx"
    d.save(path)
    assert "keep moved" in convert_file(path, front_matter=False).markdown


def test_pdf_blank_line_in_code_keeps_one_block_and_indent(tmp_path):
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    for y, t in ((100, "def f():"), (112, "    x = 1"), (136, "    return x")):
        page.insert_text((72, y), t, fontname="cour", fontsize=10)
    path = tmp_path / "cg.pdf"
    doc.save(path)
    assert "```\ndef f():\n    x = 1\n\n    return x\n```" in convert_file(path, front_matter=False).markdown


def test_docx_numbering_definitions_are_bounded(tmp_path):
    levels = "".join(f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>' for i in range(200))
    ctx = _ctx_with_numbering(tmp_path, levels)
    assert max(ctx.nums["1"]) < 9

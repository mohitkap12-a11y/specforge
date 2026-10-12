from specforge.devtools.synth import base_spec, write_docx
from specforge.ingest.audit import audit_spec
from specforge.ingest.docx import parse_docx
from specforge.ingest.pdf import parse_pdf


def test_docx_matches_pdf_semantics(demo_dir):
    d = parse_docx(str(demo_dir / "nvme-base-2.4.docx"))
    p = parse_pdf(str(demo_dir / "nvme-base-2.4.pdf"), workers=1)
    assert (d.family, d.version) == (p.family, p.version) == ("nvme-base", "2.4")
    assert [t.label for t in d.tables] == [t.label for t in p.tables]
    for td, tp in zip(d.tables, p.tables):
        assert (td.header, td.rows, td.caption) == (tp.header, tp.rows, tp.caption), td.label
    assert [s.number for s in d.sections] == [s.number for s in p.sections]
    assert [(r.text, r.strength) for r in d.requirements] == [(r.text, r.strength) for r in p.requirements]
    assert {(r.kind, r.target_family, r.target_label) for r in d.references} == {(r.kind, r.target_family, r.target_label) for r in p.references}
    issues, summary = audit_spec(d)
    assert summary["coverage_ready"], issues


def test_docx_auto_numbered_headings(tmp_path):
    p = tmp_path / "auto.docx"
    write_docx(p, base_spec(), numbered_text=False)
    s = parse_docx(str(p), "nvme-base", "2.4")
    nums = [x.number for x in s.sections]
    assert nums == ["front", "1", "2", "2.1", "2.2", "2.2.1", "2.3", "2.4", "2.5"]  # replayed counters (headings had no numbers)


def test_docx_nested_and_merged_tables(tmp_path):
    p = tmp_path / "nested.docx"
    write_docx(p, base_spec(), nested_demo=True)
    s = parse_docx(str(p), "nvme-base", "2.4")
    outer = next(t for t in s.tables if t.label == "900")
    assert outer.header == ["Bits", "Description", "Notes"]
    assert outer.rows[0][0] == "07:00" and outer.rows[1][0] == ""     # vertical merge: value lives in the first cell
    inner = [t for t in s.tables if t.parent_table == "900"]
    assert len(inner) == 1 and inner[0].rows == [["01h", "Inner meaning"]] and inner[0].parent_cell == (1, 2)
    _, summary = audit_spec(s)
    assert summary["checks"]["text_accounting"]["pass"], summary


def test_docx_continued_tables_are_stitched(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("5 Admin", 1)
    for cont in (False, True):
        d.add_paragraph("Figure 7: Big table" + (" (continued)" if cont else ""), style="Caption")
        t = d.add_table(rows=3, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text = "Value", "Description"
        base = 3 if cont else 0
        for i in (1, 2):
            t.cell(i, 0).text, t.cell(i, 1).text = f"{base + i:02d}h", f"row {base + i}"
    d.save(tmp_path / "c.docx")
    s = parse_docx(str(tmp_path / "c.docx"), "nvme-base", "2.4")
    assert len(s.tables) == 1 and [r[0] for r in s.tables[0].rows] == ["01h", "02h", "04h", "05h"]

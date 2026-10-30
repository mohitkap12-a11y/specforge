import time

import pytest
from fastapi.testclient import TestClient


@pytest.fixture()
def client(ws):
    from specforge.api.app import create_app

    with TestClient(create_app()) as c:
        yield c


def wait(client, job, timeout=120):
    t = time.time()
    while time.time() - t < timeout:
        j = client.get(f"/api/jobs/{job}").json()
        if j["state"] in ("done", "failed", "cancelled"):
            assert j["state"] == "done", j
            return j["result"]
        time.sleep(0.1)
    raise TimeoutError


def test_end_to_end_via_api(client, demo_dir):
    assert client.get("/api/health").json()["ok"]
    files = [("files", (n, open(demo_dir / n, "rb"), "application/pdf")) for n in ("nvme-base-2.4.pdf", "nvme-nvm-cs-1.3.pdf", "nvme-mi-2.0.pdf")]
    docs = client.post("/api/specs/upload", files=files).json()
    assert [d["family"] for d in docs] == ["nvme-base", "nvme-nvm-cs", "nvme-mi"]
    assert client.post("/api/specs/upload", files=[("files", ("x.txt", b"hi", "text/plain"))]).status_code == 400
    for d in docs:
        wait(client, client.post(f"/api/specs/{d['id']}/parse").json()["job"])
    listing = client.get("/api/specs").json()
    assert all(d["runs"][0]["active"] and d["runs"][0]["audit"]["coverage_ready"] for d in listing)
    run = listing[0]["runs"][0]["id"]
    tree = client.get(f"/api/runs/{run}/tree").json()
    sec = client.get(f"/api/runs/{run}/section", params={"key": next(t["key"] for t in tree if t["number"] == "5.2")}).json()
    assert [t["number"] for t in sec["tables"]] == ["101", "102", "103"] and sec["requirements"]
    # reparse produces a second, inactive run
    wait(client, client.post(f"/api/specs/{docs[0]['id']}/parse").json()["job"])
    runs = client.get("/api/specs").json()[0]["runs"]
    assert [r["active"] for r in runs] == [0, 1]
    assert client.get(f"/api/runs/{runs[1]['id']}/diff/{runs[0]['id']}").json()["summary"]["changed"] == 0

    sid = client.post("/api/sets", json={"name": "demo", "document_ids": [d["id"] for d in docs]}).json()["id"]
    assert client.get(f"/api/sets/{sid}/references").json()["summary"]["broken"] == 0
    cat = wait(client, client.post(f"/api/sets/{sid}/catalog/build").json()["job"])
    assert cat["items"] == 195 and "templates" not in cat           # the product no longer auto-generates single-command scenarios
    from seed import generate_templates

    from specforge.core.db import connect
    assert generate_templates(connect(), sid)["generated"] == 152      # test seed: gives the coverage engine scenarios to match

    src = client.post("/api/cmddb", json={"path": str(demo_dir / "cmds-a.db")}).json()["id"]
    assert client.get("/api/cmddb/suggest", params={"path": str(demo_dir / "cmds-a.db")}).json()["mapping"]["table"] == "commands"
    assert wait(client, client.post(f"/api/cmddb/{src}/import").json()["job"]) is not None
    assert client.get(f"/api/cmddb/{src}/commands", params={"opcode": 6, "queue": "admin"}).json()[0]["dwords"]["0"] & 0xFF == 6
    cov = wait(client, client.post("/api/coverage/run", json={"set_id": sid, "source_id": src}).json()["job"])["run_id"]
    summ = client.get(f"/api/coverage/{cov}").json()["summary"]["overall"]
    assert summ["covered"] == 9 and summ["stimulus_covered"] == 0
    items = client.get(f"/api/coverage/{cov}/items", params={"command": "cmd:admin:06", "state": "not_covered"}).json()
    assert items["total"] > 0
    ev = client.get(f"/api/coverage/{cov}/evidence", params={"item": "cmd:admin:06/dw10[07:00]=01h"}).json()
    assert ev[0]["result"] == "matched" and ev[0]["match_count"] == 3 and len(ev[0]["evidence"]) == 3  # seq2, seq3 and the reserved-bit command (CNS still 01h)
    for fmt, magic in (("xlsx", b"PK"), ("json", b"{"), ("html", b"<!doctype"), ("csv", b"key,")):
        r = client.get(f"/api/coverage/{cov}/export", params={"format": fmt})
        assert r.status_code == 200 and r.content.startswith(magic), fmt
    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(client.get(f"/api/coverage/{cov}/export", params={"format": "xlsx"}).content))
    assert {"Summary", "By command", "Gaps", "Scenarios", "Evidence", "Unknown commands", "Specs"} <= set(wb.sheetnames)

    # scenario review rules: an invalid predicate cannot be edited in or approved
    from specforge.core.db import connect

    c2 = connect()
    c2.execute("INSERT INTO scenario(set_id,sid,origin,title,observability,predicate_json,state) VALUES(?,?,?,?,?,?, 'draft')",
               (sid, "M:edit", "manual", "manual", "submission", "{}"))
    c2.commit()
    manual = c2.execute("SELECT id FROM scenario WHERE sid='M:edit'").fetchone()[0]
    bad = {"steps": [{"command": "cmd:admin:06", "interface": "pcie", "queue": "admin", "opcode": 99, "fields": []}]}
    assert client.patch(f"/api/scenarios/{manual}", json={"predicate": bad}).status_code == 422
    assert client.post("/api/scenarios/bulk", json={"ids": [manual], "state": "approved"}).json()["refused"][0]["id"] == manual


def test_settings_and_provider_errors(client):
    s = client.put("/api/settings", json={"tasks": {"generate": {"provider": "opencode", "model": "org/m"}}}).json()
    assert s["tasks"]["generate"]["model"] == "org/m" and s["tasks"]["classify"]["provider"] == "ollama"   # deep-merged
    assert client.get("/api/settings").json()["tasks"]["generate"]["provider"] == "opencode"
    client.put("/api/settings", json={"ollama": {"base_url": "http://127.0.0.1:9"}})
    r = client.get("/api/providers/ollama/models").json()
    assert r["ok"] is False and "failed" in r["error"]
    assert client.get("/api/jobs/9999").status_code == 404


def test_reset_wipes_work_but_keeps_settings(client, demo_dir):
    client.put("/api/settings", json={"ollama": {"base_url": "http://example:1"}})
    docs = client.post("/api/specs/upload", files=[("files", (p.name, p.read_bytes(), "application/pdf")) for p in sorted(demo_dir.glob("nvme-base-2.4.pdf"))]).json()
    client.post("/api/sets", json={"name": "demo", "document_ids": [d["id"] for d in docs]})
    r = client.post("/api/reset")
    assert r.status_code == 200 and r.json()["deleted"]["specs"] == 1
    assert client.get("/api/specs").json() == [] and client.get("/api/sets").json() == []
    assert client.get("/api/settings").json()["ollama"]["base_url"] == "http://example:1"
    from specforge.core.config import workspace_dir
    assert not list((workspace_dir() / "specs").iterdir())

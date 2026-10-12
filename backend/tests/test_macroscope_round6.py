"""Regression tests for the sixth round of Macroscope findings (review of f413206)."""
import io
import sqlite3
import time

from fastapi.testclient import TestClient
from openpyxl import load_workbook
from test_macroscope_round2 import _cmds, _scn, _step

from specforge import cmddb, coverage, jobs, reports
from specforge.core.config import db_path
from specforge.core.db import connect
from specforge.ingest.normalize import normative_info, parse_value_range, split_sentences


def test_api_returns_422_not_500_for_missing_fields(ws, conn, parsed_set):
    from specforge.api.app import create_app

    with TestClient(create_app(), raise_server_exceptions=False) as client:
        for method, url in (("put", "/api/runs/1/tables/correction"), ("post", "/api/sets/1/references/choose"),
                            ("put", "/api/sets/1/catalog/override"), ("post", "/api/scenarios/bulk"),
                            ("post", "/api/scenarios/validate"), ("post", "/api/cmddb"), ("post", "/api/coverage/run")):
            r = getattr(client, method)(url, json={})
            assert r.status_code == 422, (url, r.status_code, r.text[:80])
        ok = client.put("/api/runs/1/tables/correction", json={"key": "nvme-base@2.4/fig:101", "header": ["Bits", "Description"], "rows": []})
        assert ok.status_code == 200


def test_value_range_separator_is_case_insensitive():
    assert parse_value_range("04h to 0Fh") == parse_value_range("04h TO 0Fh") == parse_value_range("04H To 0FH") == (4, 15)


def test_etc_can_end_a_sentence():
    s = split_sentences("Use A, B, etc. The host may retry.")
    assert len(s) == 2 and normative_info(s[0]) is None and normative_info(s[1])[0] == "optional"
    assert len(split_sentences("Items A, etc. and B shall be set.")) == 1                      # lower case continues the sentence
    assert len(split_sentences("The host sets it, e.g. the Foo field. The host may retry.")) == 2


def test_cancel_ignores_unknown_and_finished_jobs(ws):
    conn = connect()
    jobs.cancel(1)                                                         # no such job yet
    jid = jobs.submit(conn, "t", lambda c, p: "ran", db_file=str(db_path()))
    assert jid == 1
    for _ in range(100):
        if jobs.get(conn, jid)["state"] in ("done", "cancelled"):
            break
        time.sleep(0.05)
    assert jobs.get(conn, jid)["state"] == "done"                          # the earlier cancel did not poison the later job
    jobs.cancel(jid)                                                       # already finished
    assert (str(db_path()), jid) not in jobs._cancel


def test_report_and_evidence_come_from_the_run_not_from_live_tables(ws, conn, parsed_set, demo_dir):
    conn.execute("INSERT INTO scenario(set_id,sid,origin,kind,title,observability,predicate_json,state) "
                 "VALUES(1,'D:1','ai','requirement','a draft','submission','{}','draft')")
    conn.commit()
    src = cmddb.register_source(conn, str(demo_dir / "cmds-a.db")); cmddb.import_source(conn, src)
    run = coverage.run_coverage(conn, 1, src)
    before = reports.build_dataset(conn, run)
    # now change everything the report used to read live
    conn.execute("UPDATE scenario SET title='CHANGED', state='rejected', predicate_json='{}'")
    conn.execute("DELETE FROM scenario WHERE id IN (SELECT id FROM scenario ORDER BY id DESC LIMIT 5)")
    conn.execute("DELETE FROM scenario_item"); conn.execute("DELETE FROM spec_set_member"); conn.execute("DELETE FROM parse_issue")
    conn.execute("UPDATE spec_document SET version='9.9'"); conn.commit()
    after = reports.build_dataset(conn, run)
    assert after == before
    assert any(s["sid"] == "D:1" and s["state"] == "draft" and s["result"] == "not_evaluated" for s in before["scenarios"])
    assert len(before["scenarios"]) == 187 and {s["family"] for s in before["specs"]} == {"nvme-base", "nvme-nvm-cs", "nvme-mi"}
    from specforge.api.app import create_app

    with TestClient(create_app()) as client:
        ev = client.get(f"/api/coverage/{run}/evidence", params={"item": "cmd:admin:06"}).json()
    assert len(ev) == 1 and ev[0]["result"] == "matched" and ev[0]["title"].startswith("Issue Identify")      # not 'CHANGED'


def test_evaluation_errors_are_visible_in_the_evidence_sheet(conn, parsed_set, tmp_path, monkeypatch):
    _cmds(tmp_path / "c.db", [(9, 1)] * 60 + [(10, 2)] * 60)
    src = cmddb.register_source(conn, str(tmp_path / "c.db")); cmddb.import_source(conn, src)
    _scn(conn, 61, [_step(9, capture="fid"), _step(9), _step(10, var="fid")], order="any")
    monkeypatch.setattr(coverage, "MAX_SEARCH_STEPS", 500)
    run = coverage.run_coverage(conn, 1, src)
    wb = load_workbook(io.BytesIO(reports.to_xlsx(reports.build_dataset(conn, run))))
    ws = wb["Evidence"]
    header = [c.value for c in ws[1]]
    assert header[-1] == "error"
    assert any(r[-1] and "exceeded" in str(r[-1]) for r in ws.iter_rows(min_row=2, values_only=True))


def test_absent_dword_clears_an_earlier_capture(conn, parsed_set, tmp_path):
    def msg(opcode, nmd0, with_nmd1=True):
        b = bytes([0x84, 0x08, 0, 0, opcode, 0, 0, 0]) + nmd0.to_bytes(4, "little")
        return b + (bytes(4) if with_nmd1 else b"")

    con = sqlite3.connect(tmp_path / "mi.db")
    con.execute("CREATE TABLE commands(id INTEGER PRIMARY KEY, test_name TEXT, seq INTEGER, mi_msg BLOB)")
    for i, m in enumerate([msg(0, 5), msg(1, 0, with_nmd1=False), msg(2, 5)], start=1):
        con.execute("INSERT INTO commands VALUES(?,?,?,?)", (i, "t", i, m))
    con.commit(); con.close()
    src = cmddb.register_source(conn, str(tmp_path / "mi.db")); cmddb.import_source(conn, src)

    def step(op, **kw):
        return {"interface": "mi", "queue": "mi", "cs": None, "opcode": op, "fields": kw.get("fields", []), **({"capture": kw["capture"]} if "capture" in kw else {})}

    s = _scn(conn, 71, [step(0, capture={"x": {"dw": "nmd0", "hi": 7, "lo": 0}}),
                        step(1, capture={"x": {"dw": "nmd1", "hi": 7, "lo": 0}}),          # nmd1 is absent in this message
                        step(2, fields=[{"dw": "nmd0", "hi": 7, "lo": 0, "op": "eq", "value": {"var": "x"}}])])
    run = coverage.run_coverage(conn, 1, src)
    res = conn.execute("SELECT state FROM coverage_result WHERE run_id=? AND scenario_id=?", (run, s)).fetchone()[0]
    assert res == "not_matched"        # x must not silently keep the value 5 captured by the first step

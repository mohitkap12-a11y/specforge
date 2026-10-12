"""Persistence for spec documents, parse runs, IR nodes; reparse, diff, corrections, IR loading."""
from __future__ import annotations

import hashlib
import json
import shutil
import sqlite3
from pathlib import Path
from typing import Any, Callable

from ..core.config import PARSER_VERSION, workspace_dir
from ..core.db import jloads, row, rows
from ..core.hashing import sha256_file, sha256_text, stable_hash
from .audit import audit_spec
from .ir import ParsedSpec
from .normalize import norm_title

ProgressFn = Callable[[str, int, int], None]


def _fmt_of(path: str | Path) -> str:
    suf = Path(path).suffix.lower()
    if suf == ".pdf":
        return "pdf"
    if suf in (".docx", ".dotx"):
        return "docx"
    raise ValueError(f"Unsupported spec format: {suf}")


def detect_document(path: str | Path, filename: str | None = None) -> dict[str, Any]:
    """Cheap identity detection used by the upload step (family / version / pages)."""
    from .normalize import detect_identity

    fmt = _fmt_of(path)
    if fmt == "pdf":
        import pymupdf

        doc = pymupdf.open(str(path))
        text = "\f".join(doc[i].get_text("text") for i in range(min(3, doc.page_count)))
        pages = doc.page_count
        title = (doc.metadata or {}).get("title") or ""
        doc.close()
    else:
        import docx

        d = docx.Document(str(path))
        text = "\f".join(p.text for p in d.paragraphs[:80])
        pages, title = 0, d.core_properties.title or ""
    fam, ver = detect_identity(text, filename or Path(path).name)
    return {"family": fam, "version": ver, "pages": pages, "title": title, "format": fmt}


def register_document(conn: sqlite3.Connection, src_path: str | Path, family: str | None = None,
                      version: str | None = None, filename: str | None = None) -> dict:
    """Copy the file into the workspace (content-addressed) and create/return the spec_document row."""
    src = Path(src_path)
    sha = sha256_file(src)
    existing = row(conn, "SELECT * FROM spec_document WHERE sha256=?", (sha,))
    if existing:
        return existing
    info = detect_document(src, filename)
    family = family or info["family"] or "unknown"
    version = version or info["version"] or "0"
    dest = workspace_dir() / "specs" / f"{sha[:16]}{src.suffix.lower()}"
    if not dest.exists():
        shutil.copyfile(src, dest)
    cur = conn.execute(
        "INSERT INTO spec_document(family,version,title,sha256,path,filename,format,pages) VALUES(?,?,?,?,?,?,?,?)",
        (family, version, info["title"], sha, str(dest), filename or src.name, info["format"], info["pages"]))
    conn.commit()
    return row(conn, "SELECT * FROM spec_document WHERE id=?", (cur.lastrowid,))  # type: ignore[return-value]


def parse_config(settings: dict) -> dict:
    p = settings.get("parser", {})
    return {"table_strategy": p.get("table_strategy", "auto")}


def run_parse(conn: sqlite3.Connection, document_id: int, settings: dict | None = None, workers: int | None = None,
              progress: ProgressFn | None = None, activate: bool | None = None) -> int:
    """Parse a registered document into a new parse_run. Returns the run id."""
    import os

    from ..core.config import load_settings
    from .docx import parse_docx
    from .pdf import parse_pdf

    settings = settings or load_settings(conn)
    doc = row(conn, "SELECT * FROM spec_document WHERE id=?", (document_id,))
    if not doc:
        raise KeyError(f"document {document_id}")
    cfg = parse_config(settings)
    cur = conn.execute("INSERT INTO parse_run(document_id,parser_version,config_hash,status) VALUES(?,?,?, 'running')",
                       (document_id, PARSER_VERSION, stable_hash(cfg)))
    run_id = cur.lastrowid
    conn.commit()
    try:
        w = workers if workers is not None else settings.get("parser", {}).get("workers", 0)
        w = w or max(1, (os.cpu_count() or 2) - 1)
        if doc["format"] == "pdf":
            spec = parse_pdf(doc["path"], doc["family"], doc["version"], workers=w, progress=progress)
        else:
            spec = parse_docx(doc["path"], doc["family"], doc["version"],
                              progress=(lambda s, a, b: progress(s, a, b)) if progress else None)
        spec.family, spec.version = doc["family"], doc["version"]
        apply_corrections(conn, spec)
        issues, summary = audit_spec(spec)
        _store_ir(conn, run_id, spec, issues)
        stats = {"pages": spec.pages, "sections": len(spec.sections) - 1, "tables": len(spec.tables),
                 "requirements": len(spec.requirements), "references": len(spec.references),
                 "excluded_lines": len(spec.excluded), **spec.stats}
        make_active = activate if activate is not None else not row(
            conn, "SELECT 1 AS x FROM parse_run WHERE document_id=? AND active=1", (document_id,))
        conn.execute("UPDATE parse_run SET status='done', stats_json=?, audit_json=? WHERE id=?",
                     (json.dumps(stats), json.dumps(summary), run_id))
        conn.commit()
        if make_active:
            activate_run(conn, run_id)
    except Exception as e:  # noqa: BLE001
        conn.execute("UPDATE parse_run SET status='failed', error=? WHERE id=?", (repr(e), run_id))
        conn.commit()
        raise
    return run_id  # type: ignore[return-value]


def activate_run(conn: sqlite3.Connection, run_id: int) -> None:
    r = row(conn, "SELECT document_id, status FROM parse_run WHERE id=?", (run_id,))
    if not r:
        raise KeyError(run_id)
    if r["status"] != "done":
        raise ValueError(f"parse run {run_id} is '{r['status']}'; only finished runs can be activated")
    conn.execute("UPDATE parse_run SET active=0 WHERE document_id=?", (r["document_id"],))
    conn.execute("UPDATE parse_run SET active=1 WHERE id=?", (run_id,))
    conn.commit()


def _store_ir(conn: sqlite3.Connection, run_id: int, spec: ParsedSpec, issues: list[dict]) -> None:
    prefix = spec.prefix
    nodes: list[tuple] = []
    ordn = 0
    for s in spec.sections:
        ordn += 1
        key = f"{prefix}/sec:{s.number}"
        parent = f"{prefix}/sec:{s.parent}" if s.parent else None
        nodes.append((run_id, key, "section", parent, ordn, s.number, s.title, "\n\n".join(s.paragraphs), s.page_from,
                      s.page_to, json.dumps({"level": s.level})))
    for t in spec.tables:
        ordn += 1
        key = f"{prefix}/fig:{t.label}"
        parent = f"{prefix}/fig:{t.parent_table}" if t.parent_table else (f"{prefix}/sec:{t.section}" if t.section else f"{prefix}/sec:front")
        nodes.append((run_id, key, "table", parent, ordn, t.label, t.caption, None, t.page_from, t.page_to,
                      json.dumps({"header": t.header, "rows": t.rows, "captioned": t.captioned, "extractor": t.extractor,
                                  "quality": t.quality, "section": t.section, "parent_cell": t.parent_cell,
                                  "continued_pages": t.continued_pages})))
    seen: dict[str, int] = {}
    for r in spec.requirements:
        ordn += 1
        h = hashlib.sha1(norm_title(r.text).encode()).hexdigest()[:8]
        base = f"{prefix}/{('fig:' + r.table) if r.table else ('sec:' + (r.section or 'front'))}/req:{h}"
        seen[base] = seen.get(base, 0) + 1
        key = base if seen[base] == 1 else f"{base}.{seen[base]}"
        parent = f"{prefix}/fig:{r.table}" if r.table else f"{prefix}/sec:{r.section or 'front'}"
        nodes.append((run_id, key, "requirement", parent, ordn, None, None, r.text, r.page, r.page,
                      json.dumps({"strength": r.strength, "word": r.normative_word, "section": r.section, "table": r.table})))
    conn.executemany("INSERT INTO ir_node(run_id,key,kind,parent_key,ord,number,title,text,page_from,page_to,attrs_json)"
                     " VALUES(?,?,?,?,?,?,?,?,?,?,?)", nodes)
    refs = [(run_id, f"{prefix}/{r.src_key}", r.raw, r.kind, r.target_family, r.target_label, i)
            for i, r in enumerate(spec.references)]
    conn.executemany("INSERT INTO ir_ref(run_id,src_key,raw,kind,target_family,target_label,ord) VALUES(?,?,?,?,?,?,?)", refs)
    conn.executemany("INSERT INTO parse_issue(run_id,severity,kind,node_key,page,message) VALUES(?,?,?,?,?,?)",
                     [(run_id, i["severity"], i["kind"], (f"{prefix}/{i['node_key']}" if i.get("node_key") else None),
                       i.get("page"), i["message"]) for i in issues])
    conn.commit()


# ----------------------------------------------------------------------------- corrections

def set_correction(conn: sqlite3.Connection, family: str, version: str, stable_key: str, kind: str, payload: dict) -> None:
    conn.execute("INSERT INTO correction(family,version,stable_key,kind,payload_json) VALUES(?,?,?,?,?) "
                 "ON CONFLICT(family,version,stable_key,kind) DO UPDATE SET payload_json=excluded.payload_json",
                 (family, version, stable_key, kind, json.dumps(payload)))
    conn.commit()


def apply_corrections(conn: sqlite3.Connection, spec: ParsedSpec) -> None:
    """Re-apply user table corrections; corrections whose target vanished are reported, never dropped."""
    for c in rows(conn, "SELECT * FROM correction WHERE family=? AND version=?", (spec.family, spec.version)):
        label = c["stable_key"].split("fig:", 1)[-1]
        t = next((x for x in spec.tables if x.label == label), None)
        if t is None:
            spec.notes.append({"severity": "warning", "kind": "correction_orphaned", "node_key": c["stable_key"], "page": None,
                               "message": f"Correction for {c['stable_key']} no longer matches any table after reparse"})
            continue
        payload = json.loads(c["payload_json"])
        if c["kind"] == "table_replace":
            t.header = payload.get("header", t.header)
            t.rows = payload.get("rows", t.rows)
            t.quality = 1.0
            spec.notes.append({"severity": "info", "kind": "correction_applied", "node_key": f"fig:{t.label}",
                               "page": t.page_from, "message": "User correction applied to this table"})
    # corrected tables changed text; rebuild requirements and references
    from .semantics import build_semantics
    build_semantics(spec)
    # drop stale quality warnings for corrected tables
    corrected = {c["stable_key"].split("fig:", 1)[-1] for c in rows(
        conn, "SELECT stable_key FROM correction WHERE family=? AND version=? AND kind='table_replace'",
        (spec.family, spec.version))}
    spec.notes = [n for n in spec.notes if not (n["kind"] in ("low_table_quality", "table_text_loss") and
                                                (n.get("node_key") or "").split("fig:")[-1] in corrected)]


# ----------------------------------------------------------------------------- loading & diffing

def load_ir(conn: sqlite3.Connection, run_id: int) -> dict[str, Any]:
    run = row(conn, "SELECT r.*, d.family, d.version, d.format FROM parse_run r JOIN spec_document d ON d.id=r.document_id WHERE r.id=?", (run_id,))
    if not run:
        raise KeyError(run_id)
    prefix = f"{run['family']}@{run['version']}"
    ir: dict[str, Any] = {"run_id": run_id, "family": run["family"], "version": run["version"], "prefix": prefix,
                          "sections": {}, "section_order": [], "tables": {}, "requirements": [], "refs": []}
    for n in rows(conn, "SELECT * FROM ir_node WHERE run_id=? ORDER BY ord", (run_id,)):
        a = jloads(n["attrs_json"], {})
        if n["kind"] == "section":
            ir["sections"][n["number"]] = {"key": n["key"], "number": n["number"], "title": n["title"], "text": n["text"] or "",
                                           "parent": n["parent_key"].split("sec:", 1)[-1] if n["parent_key"] else None,
                                           "page_from": n["page_from"], "page_to": n["page_to"], "level": a.get("level", 0)}
            ir["section_order"].append(n["number"])
        elif n["kind"] == "table":
            ir["tables"][n["number"]] = {"key": n["key"], "label": n["number"], "caption": n["title"] or "", "page_from": n["page_from"],
                                         "parent_key": n["parent_key"],
                                         "page_to": n["page_to"], **a}
        else:
            ir["requirements"].append({"key": n["key"], "text": n["text"], "page": n["page_from"], **a})
    ir["refs"] = rows(conn, "SELECT * FROM ir_ref WHERE run_id=? ORDER BY ord", (run_id,))
    return ir


def diff_runs(conn: sqlite3.Connection, a: int, b: int) -> dict[str, Any]:
    def snap(run_id: int) -> dict[str, str]:
        out = {}
        for n in rows(conn, "SELECT key,kind,title,text,parent_key,attrs_json FROM ir_node WHERE run_id=?", (run_id,)):
            a_ = jloads(n["attrs_json"], {})
            body = a_.get("rows") if n["kind"] == "table" else n["text"]
            structure = {k: a_.get(k) for k in ("section", "parent_cell", "captioned", "level", "strength", "table")}
            out[n["key"]] = sha256_text(json.dumps([n["title"], body, a_.get("header"), n["parent_key"], structure], default=str))
        return out

    sa, sb = snap(a), snap(b)
    added = sorted(set(sb) - set(sa))
    removed = sorted(set(sa) - set(sb))
    changed = sorted(k for k in set(sa) & set(sb) if sa[k] != sb[k])
    return {"from": a, "to": b, "added": added, "removed": removed, "changed": changed,
            "summary": {"added": len(added), "removed": len(removed), "changed": len(changed), "unchanged": len(set(sa) & set(sb)) - len(changed)}}


def reparse_available(conn: sqlite3.Connection) -> list[int]:
    """Documents whose active run was produced by an older parser version."""
    return [r["document_id"] for r in rows(
        conn, "SELECT document_id FROM parse_run WHERE active=1 AND parser_version<>?", (PARSER_VERSION,))]

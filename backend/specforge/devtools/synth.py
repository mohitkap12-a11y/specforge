"""Synthetic NVMe-like specification fixtures (PDF + DOCX) and a synthetic command database.

These are NOT the real NVMe specs.  They imitate the structures the parser must handle: numbered
sections with bookmarks, 'Figure N: ...' captions above ruled tables, multi-page tables with repeated
headers, cross-spec references, normative language, nested/merged tables (DOCX), headers/footers.
Opcode values follow the NVMe numbering so decoded demo data looks realistic.
"""
from __future__ import annotations

import sqlite3
import struct
from pathlib import Path

Block = tuple

# ----------------------------------------------------------------------------- content model

def _lid_rows(n: int = 62) -> list[list[str]]:
    known = [("01h", "M", "Error Information"), ("02h", "M", "SMART / Health Information"),
             ("03h", "M", "Firmware Slot Information"), ("04h", "O", "Changed Namespace List"),
             ("05h", "O", "Commands Supported and Effects"), ("06h", "O", "Device Self-test")]
    rows = [list(r) for r in known]
    for i in range(7, n):
        rows.append([f"{i:02X}h", "O", f"Synthetic Log Page {i:02X}h"])
    return rows


def base_spec(lid_rows: int = 62) -> list[Block]:
    return [
        ("title", "NVM Express® Base Specification", "Revision 2.4"),
        ("h", 1, "1 Introduction"),
        ("p", "This synthetic document imitates the layout of the NVM Express Base Specification. "
              "A controller shall process each command in a submission queue entry. "
              "A host may submit multiple commands to an Admin Queue."),
        ("h", 1, "5 Admin Command Set"),
        ("h", 2, "5.1 Admin Command Overview"),
        ("p", "Figure 100 lists the opcodes for Admin commands. A controller shall abort a command with an unsupported opcode "
              "with status Invalid Command Opcode."),
        ("table", "Figure 100: Opcodes for Admin Commands", ["Opcode", "O/M", "Command"],
         [["01h", "M", "Create I/O Submission Queue"], ["02h", "M", "Get Log Page"], ["05h", "M", "Create I/O Completion Queue"],
          ["06h", "M", "Identify"], ["08h", "M", "Abort"], ["09h", "M", "Set Features"], ["0Ah", "M", "Get Features"],
          ["0Dh", "O", "Namespace Management"], ["15h", "O", "Namespace Attachment"]], None),
        ("h", 2, "5.2 Identify command"),
        ("p", "The Identify command returns a data structure to the host. The controller shall abort the command with status "
              "Invalid Field in Command if a reserved field is set to a non-zero value. The CNS field is defined in Figure 103. "
              "Refer to section 3.1 of the NVM Express NVM Command Set Specification for I/O namespace data structures."),
        ("table", "Figure 101: Identify – Command Dword 10", ["Bits", "Description"],
         [["31:16", "Controller Identifier (CNTID): This field specifies the controller identifier used as part of some CNS values."],
          ["15:08", "Reserved"],
          ["07:00", "Controller or Namespace Structure (CNS): This field specifies the information to be returned to the host. "
                    "Refer to Figure 103 for the defined values. The host shall not specify a reserved value."]], None),
        ("table", "Figure 102: Identify – Command Dword 11", ["Bits", "Description"],
         [["31:24", "Command Set Identifier (CSI): This field specifies the I/O Command Set to be returned."],
          ["23:16", "Reserved"],
          ["15:00", "CNS Specific Identifier (CNSSPECID): This field is used by some CNS values."]], None),
        ("table", "Figure 103: Identify – CNS Values", ["CNS Value", "O/M", "Description", "Reference Section"],
         [["00h", "M", "Identify Namespace data structure for the specified NSID", "5.2.1"],
          ["01h", "M", "Identify Controller data structure for the controller processing the command", "5.2.2"],
          ["02h", "O", "Active Namespace ID list", "5.2.3"],
          ["03h", "O", "Namespace Identification Descriptor list", "5.2.4"],
          ["04h to 0Fh", "", "Reserved", ""],
          ["10h", "O", "Allocated Namespace ID list", "5.2.5"]], None),
        ("h", 3, "5.2.1 Identify Namespace data structure"),
        ("p", "This data structure describes a namespace. Bytes 07:00 contain the namespace size. The controller shall "
              "return zeros for unused bytes."),
        ("table", "Figure 104: Identify Namespace Data Structure", ["Bytes", "O/M", "Description"],
         [["07:00", "M", "Namespace Size (NSZE): This field indicates the total size of the namespace in logical blocks."],
          ["15:08", "M", "Namespace Capacity (NCAP): This field indicates the maximum number of logical blocks that may be allocated."]], None),
        ("h", 2, "5.3 Get Log Page command"),
        ("p", "The Get Log Page command returns a log page. The host should set the Retain Asynchronous Event bit appropriately."),
        ("table", "Figure 120: Get Log Page – Command Dword 10", ["Bits", "Description"],
         [["31:16", "Number of Dwords Lower (NUMDL): This field specifies the number of dwords to return, 0's based."],
          ["15", "Retain Asynchronous Event (RAE): When set to 1, the controller shall retain the asynchronous event."],
          ["14:08", "Log Specific Field (LSP): This field is log page specific."],
          ["07:00", "Log Page Identifier (LID): This field specifies the identifier of the log page. Refer to Figure 121."]], None),
        ("table", "Figure 121: Get Log Page – Log Page Identifiers", ["Log Page Identifier", "O/M", "Log Page Name"],
         _lid_rows(lid_rows), None),
        ("h", 2, "5.4 Set Features command"),
        ("p", "The Set Features command configures a feature. If the Feature Identifier is not supported, the controller shall "
              "abort the command with status Invalid Field in Command."),
        ("table", "Figure 130: Set Features – Command Dword 10", ["Bits", "Description"],
         [["31", "Save (SV): When set to 1, the controller shall save the attribute."],
          ["30:08", "Reserved"],
          ["07:00", "Feature Identifier (FID): This field specifies the identifier of the feature. Refer to Figure 131."]], None),
        ("table", "Figure 131: Feature Identifiers", ["Feature Identifier", "O/M", "Feature Name"],
         [["01h", "M", "Arbitration"], ["02h", "M", "Power Management"], ["04h", "M", "Temperature Threshold"],
          ["07h", "M", "Number of Queues"], ["0Bh", "M", "Asynchronous Event Configuration"]], None),
        ("h", 2, "5.5 Get Features command"),
        ("p", "The Get Features command retrieves the attributes of a feature."),
        ("table", "Figure 140: Get Features – Command Dword 10", ["Bits", "Description"],
         [["31:11", "Reserved"],
          ["10:08", "Select (SEL): This field specifies which value of the attributes to return. Refer to Figure 141."],
          ["07:00", "Feature Identifier (FID): This field specifies the identifier of the feature that attributes are being retrieved for."]], None),
        ("table", "Figure 141: Get Features – Select Values", ["Value", "Description"],
         [["000b", "Current"], ["001b", "Default"], ["010b", "Saved"], ["011b", "Supported capabilities"],
          ["100b to 111b", "Reserved"]], None),
    ]


def nvm_cs_spec() -> list[Block]:
    return [
        ("title", "NVM Express® NVM Command Set Specification", "Revision 1.3"),
        ("h", 1, "1 Introduction"),
        ("p", "This synthetic document imitates the NVM Command Set Specification. It extends the NVM Express Base Specification."),
        ("h", 1, "3 I/O Commands"),
        ("h", 2, "3.1 Command Set Overview"),
        ("p", "Figure 20 lists the opcodes for the NVM command set. Admin commands are defined in section 5.1 of the "
              "NVM Express Base Specification. A controller shall abort an I/O command whose opcode is not supported."),
        ("table", "Figure 20: Opcodes for NVM Command Set", ["Opcode", "O/M", "Command"],
         [["00h", "M", "Flush"], ["01h", "M", "Write"], ["02h", "M", "Read"], ["08h", "O", "Write Zeroes"],
          ["09h", "O", "Dataset Management"]], None),
        ("h", 2, "3.2 Read command"),
        ("p", "The Read command reads data and metadata from the namespace. If the Limited Retry bit is set, the controller "
              "should apply limited error recovery. The controller shall abort the command if the LBA range exceeds the namespace capacity."),
        ("table", "Figure 30: Read – Command Dword 10", ["Bits", "Description"],
         [["31:00", "Starting LBA (SLBA): This field specifies the 64-bit address of the first logical block, least significant 32 bits."]], None),
        ("table", "Figure 31: Read – Command Dword 12", ["Bits", "Description"],
         [["31", "Limited Retry (LR): When set to 1, the controller should apply limited retry efforts."],
          ["30", "Force Unit Access (FUA): When set to 1, the controller shall read from the non-volatile media."],
          ["29:26", "Protection Information Check (PRINFO): This field specifies the protection information action."],
          ["25:16", "Reserved"],
          ["15:00", "Number of Logical Blocks (NLB): This field indicates the number of logical blocks to be read. This is a 0's based value."]], None),
        ("h", 2, "3.3 Write command"),
        ("p", "The Write command writes data and metadata to the namespace."),
        ("table", "Figure 40: Write – Command Dword 12", ["Bits", "Description"],
         [["31", "Limited Retry (LR): When set to 1, the controller should apply limited retry efforts."],
          ["30", "Force Unit Access (FUA): When set to 1, the controller shall write data to non-volatile media before indicating completion."],
          ["29:16", "Reserved"],
          ["15:00", "Number of Logical Blocks (NLB): This field indicates the number of logical blocks to be written. This is a 0's based value."]], None),
    ]


def mi_spec() -> list[Block]:
    return [
        ("title", "NVM Express® Management Interface Specification", "Revision 2.0"),
        ("h", 1, "1 Introduction"),
        ("p", "This synthetic document imitates the NVM Express Management Interface Specification. NVMe Admin commands may be "
              "tunnelled using the NVMe Admin Command message as defined in section 5.2 of the NVM Express Base Specification."),
        ("h", 1, "5 NVMe-MI Commands"),
        ("h", 2, "5.1 Command Overview"),
        ("p", "Figure 30 lists the NVMe-MI command opcodes. A Management Endpoint shall return an error response if the opcode is not supported."),
        ("table", "Figure 30: NVMe-MI Command Opcodes", ["Opcode", "O/M", "Command"],
         [["00h", "M", "Read NVMe-MI Data Structure"], ["01h", "M", "NVM Subsystem Health Status Poll"],
          ["02h", "M", "Controller Health Status Poll"], ["03h", "M", "Configuration Set"], ["04h", "M", "Configuration Get"]], None),
        ("h", 2, "5.2 Read NVMe-MI Data Structure"),
        ("p", "This command returns an NVMe-MI data structure. The Data Structure Type is defined in Figure 32."),
        ("table", "Figure 31: Read NVMe-MI Data Structure – Request Dword 0 (NMD0)", ["Bits", "Description"],
         [["31:24", "Data Structure Type (DTYP): This field specifies the data structure to read. Refer to Figure 32."],
          ["23:16", "Port Identifier (PORTID): This field specifies the port."],
          ["15:00", "Controller Identifier (CTRLID): This field specifies the controller."]], None),
        ("table", "Figure 32: NVMe-MI Data Structure Types", ["Value", "Description"],
         [["00h", "NVM Subsystem Information"], ["01h", "Port Information"], ["02h", "Controller List"],
          ["03h", "Controller Information"], ["04h to FFh", "Reserved"]], None),
    ]


SPECS = {"nvme-base": ("2.4", base_spec), "nvme-nvm-cs": ("1.3", nvm_cs_spec), "nvme-mi": ("2.0", mi_spec)}

# ----------------------------------------------------------------------------- PDF writer

def write_pdf(path: str | Path, blocks: list[Block], header: str = "NVM Express Synthetic Spec",
              footer_prefix: str = "Page") -> Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import BaseDocTemplate, Frame, PageBreak, PageTemplate, Paragraph, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=styles["BodyText"], fontName="Helvetica", fontSize=10, leading=13)
    cell = ParagraphStyle("c", parent=body, fontSize=8.5, leading=10.5)
    cellb = ParagraphStyle("cb", parent=cell, fontName="Helvetica-Bold")
    cap = ParagraphStyle("cap", parent=body, fontName="Helvetica-Bold", fontSize=10, spaceBefore=8, spaceAfter=4)
    heads = {i: ParagraphStyle(f"h{i}", parent=body, fontName="Helvetica-Bold", fontSize=15 - 2 * i, leading=18 - 2 * i,
                               spaceBefore=10, spaceAfter=5) for i in (1, 2, 3, 4)}

    class Doc(BaseDocTemplate):
        def afterFlowable(self, fl):
            if getattr(fl, "_h", None):
                lvl, text = fl._h
                key = f"h{id(fl)}"
                self.canv.bookmarkPage(key)
                self.canv.addOutlineEntry(text, key, level=lvl - 1, closed=False)

    def deco(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(inch, letter[1] - 0.55 * inch, header)
        canvas.drawRightString(letter[0] - inch, 0.5 * inch, f"{footer_prefix} {doc.page}")
        canvas.drawString(inch, 0.5 * inch, "Copyright synthetic fixture")
        canvas.restoreState()

    doc = Doc(str(path), pagesize=letter, leftMargin=inch, rightMargin=inch, topMargin=inch, bottomMargin=inch)
    doc.addPageTemplates([PageTemplate(id="p", frames=[Frame(inch, inch, letter[0] - 2 * inch, letter[1] - 2 * inch, id="f")],
                                       onPage=deco)])
    story = []
    W = letter[0] - 2 * inch
    for b in blocks:
        if b[0] == "title":
            story += [Spacer(1, 1.2 * inch), Paragraph(b[1], ParagraphStyle("t", parent=body, fontSize=24, leading=28)),
                      Spacer(1, 12), Paragraph(b[2], ParagraphStyle("t2", parent=body, fontSize=14)), PageBreak()]
        elif b[0] == "h":
            p = Paragraph(b[2], heads[min(b[1], 4)])
            p._h = (b[1], b[2])
            story.append(p)
        elif b[0] == "p":
            story.append(Paragraph(b[1], body))
        elif b[0] == "pagebreak":
            story.append(PageBreak())
        elif b[0] == "table":
            _, caption, header_row, rows, widths = b
            if caption:
                story.append(Paragraph(caption, cap))
            data = [[Paragraph(c, cellb) for c in header_row]] + [[Paragraph(c, cell) for c in r] for r in rows]
            ncol = len(header_row)
            if widths is None:
                widths = [0.9 * inch] * (ncol - 1) + [W - 0.9 * inch * (ncol - 1)] if ncol >= 2 else [W]
                if ncol == 3:
                    widths = [1.1 * inch, 0.6 * inch, W - 1.7 * inch]
                if ncol == 4:
                    widths = [1.0 * inch, 0.5 * inch, W - 2.4 * inch, 0.9 * inch]
            t = Table(data, colWidths=widths, repeatRows=1)
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.6, colors.black), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dddddd"))]))
            story.append(t)
            story.append(Spacer(1, 8))
    doc.build(story)
    return Path(path)


# ----------------------------------------------------------------------------- DOCX writer

def write_docx(path: str | Path, blocks: list[Block], numbered_text: bool = True, nested_demo: bool = False) -> Path:
    import docx
    from docx.oxml.ns import qn  # noqa: F401

    d = docx.Document()
    for b in blocks:
        if b[0] == "title":
            d.core_properties.title = b[1]
            d.add_paragraph(b[1])
            d.add_paragraph(b[2])
        elif b[0] == "h":
            text = b[2]
            if not numbered_text:
                text = text.split(" ", 1)[1]
            d.add_heading(text, level=b[1])
        elif b[0] == "p":
            d.add_paragraph(b[1])
        elif b[0] == "table":
            _, caption, header_row, rows, _w = b
            d.add_paragraph(caption, style="Caption")
            t = d.add_table(rows=1 + len(rows), cols=len(header_row))
            t.style = "Table Grid"
            for j, h in enumerate(header_row):
                t.cell(0, j).text = h
            for i, r in enumerate(rows, start=1):
                for j, c in enumerate(r):
                    t.cell(i, j).text = c
    if nested_demo:
        d.add_heading("9 Nested structures", level=1)
        d.add_paragraph("The following table contains a nested table and merged cells.")
        d.add_paragraph("Figure 900: Nested Demo", style="Caption")
        t = d.add_table(rows=3, cols=3)
        t.style = "Table Grid"
        t.cell(0, 0).text, t.cell(0, 1).text, t.cell(0, 2).text = "Bits", "Description", "Notes"
        a = t.cell(1, 0).merge(t.cell(2, 0))
        a.text = "07:00"
        t.cell(1, 1).text = "Outer description. The controller shall honour the inner table."
        t.cell(2, 1).text = "Second row"
        inner = t.cell(1, 2).add_table(rows=2, cols=2)
        inner.style = "Table Grid"
        inner.cell(0, 0).text, inner.cell(0, 1).text = "Value", "Meaning"
        inner.cell(1, 0).text, inner.cell(1, 1).text = "01h", "Inner meaning"
        t.cell(2, 2).text = "n/a"
    d.save(str(path))
    return Path(path)


def make_demo_specs(out_dir: str | Path, docx_too: bool = True) -> dict[str, Path]:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    paths = {}
    for fam, (ver, fn) in SPECS.items():
        p = out / f"{fam}-{ver}.pdf"
        write_pdf(p, fn(), header=f"{fam} {ver} synthetic")
        paths[fam] = p
        if docx_too:
            write_docx(out / f"{fam}-{ver}.docx", fn())
    return paths


# ----------------------------------------------------------------------------- synthetic command DB

def sqe(opcode: int, nsid: int = 0, cid: int = 1, cdw10: int = 0, cdw11: int = 0, cdw12: int = 0, cdw13: int = 0,
        cdw14: int = 0, cdw15: int = 0, fuse: int = 0, psdt: int = 0, cdw2: int = 0, cdw3: int = 0,
        mptr: int = 0, prp1: int = 0, prp2: int = 0) -> bytes:
    cdw0 = (opcode & 0xFF) | ((fuse & 3) << 8) | ((psdt & 3) << 14) | ((cid & 0xFFFF) << 16)
    return struct.pack("<16I", cdw0, nsid, cdw2, cdw3, mptr & 0xFFFFFFFF, mptr >> 32, prp1 & 0xFFFFFFFF, prp1 >> 32,
                       prp2 & 0xFFFFFFFF, prp2 >> 32, cdw10, cdw11, cdw12, cdw13, cdw14, cdw15)


def mi_msg(nmimt: int, opcode: int, nmd0: int = 0, nmd1: int = 0, ror: int = 0, csi: int = 0) -> bytes:
    """Build a message with the default (configurable) header layout used by cmddb.MI_DEFAULT_LAYOUT."""
    b0 = 0x84  # IC=1, MCTP message type 4h
    b1 = ((ror & 1) << 7) | ((nmimt & 0xF) << 3) | (csi & 1)
    hdr = bytes([b0, b1, 0, 0])
    return hdr + bytes([opcode & 0xFF, 0, 0, 0]) + struct.pack("<II", nmd0, nmd1)


def mi_admin_msg(opcode: int, nsid: int = 0, cdw10: int = 0, cdw11: int = 0, cdw12: int = 0, ctlid: int = 0) -> bytes:
    """NVMe Admin command carried in an NVMe-MI message (NMIMT=2) using cmddb.MI_DEFAULT_LAYOUT offsets."""
    from ..cmddb import MI_DEFAULT_LAYOUT as L

    buf = bytearray(64)
    buf[0] = 0x84
    buf[1] = (2 & 0xF) << 3
    buf[L["admin"]["opcode_byte"]] = opcode & 0xFF
    struct.pack_into("<H", buf, L["admin"]["ctlid_off"], ctlid)
    for idx, val in (("1", nsid), ("10", cdw10), ("11", cdw11), ("12", cdw12)):
        struct.pack_into("<I", buf, L["admin"]["dwords"][idx], val)
    return bytes(buf)


def make_command_db(path: str | Path, layout: str = "A", seed: int = 7) -> dict:
    """Create a synthetic command dump with known contents; returns the ground truth used by tests."""
    p = Path(path)
    if p.exists():
        p.unlink()
    con = sqlite3.connect(p)
    if layout == "A":
        con.execute("""CREATE TABLE commands(id INTEGER PRIMARY KEY, test_name TEXT, script TEXT, seq INTEGER, ts TEXT,
                       interface TEXT, queue TEXT, qid INTEGER, sqe BLOB, mi_msg BLOB)""")
    else:
        cols = ", ".join(f"cdw{i} INTEGER" for i in range(16))
        con.execute(f"""CREATE TABLE nvme_cmds(rowid_ INTEGER PRIMARY KEY, testcase TEXT, script TEXT, step INTEGER,
                       queue TEXT, {cols})""")
    rows = []
    truth = {"commands": []}

    def add(test, seq, iface, queue, sqe_b=None, mi_b=None, script="synthetic.py"):
        rows.append((test, script, seq, "2026-01-01T00:00:%02d" % (seq % 60), iface, queue, 0 if queue == "admin" else 1,
                     sqe_b, mi_b))
        truth["commands"].append({"test": test, "seq": seq, "iface": iface, "queue": queue})

    # test_identify: Identify CNS 00h, 01h (x2 with different CID), 02h ; reserved bits set once
    t = "test_identify"
    add(t, 1, "pcie", "admin", sqe(0x06, nsid=1, cid=1, cdw10=0x00))
    add(t, 2, "pcie", "admin", sqe(0x06, nsid=0, cid=2, cdw10=0x01))
    add(t, 3, "pcie", "admin", sqe(0x06, nsid=0, cid=3, cdw10=0x01, prp1=0x1000))   # same signature as seq 2
    add(t, 4, "pcie", "admin", sqe(0x06, nsid=0, cid=4, cdw10=0x02))
    add(t, 5, "pcie", "admin", sqe(0x06, nsid=0, cid=5, cdw10=0x0100 | 0x01))        # reserved CDW10[15:08] != 0
    # test_features: set then get FID 07h; get with SEL=1 (default)
    t = "test_features"
    add(t, 1, "pcie", "admin", sqe(0x09, cid=1, cdw10=0x07, cdw11=0x00010001))
    add(t, 2, "pcie", "admin", sqe(0x0A, cid=2, cdw10=0x07))
    add(t, 3, "pcie", "admin", sqe(0x0A, cid=3, cdw10=(1 << 8) | 0x04))              # SEL=1 FID=04h
    add(t, 4, "pcie", "admin", sqe(0x09, cid=4, cdw10=0x80000000 | 0x01))            # SV=1 FID=01h
    # test_io: read with LR and various NLB, write w/o flags
    t = "test_io"
    add(t, 1, "pcie", "io", sqe(0x02, nsid=1, cid=1, cdw10=0, cdw12=(1 << 31) | 7))
    add(t, 2, "pcie", "io", sqe(0x02, nsid=1, cid=2, cdw10=0x100, cdw12=0xFFFF))
    add(t, 3, "pcie", "io", sqe(0x01, nsid=1, cid=3, cdw10=0, cdw12=0))
    add(t, 4, "pcie", "io", sqe(0x00, nsid=1, cid=4))
    # test_log: get log page LID 02h, and an unknown vendor admin opcode
    t = "test_log"
    add(t, 1, "pcie", "admin", sqe(0x02, nsid=0xFFFFFFFF, cid=1, cdw10=(0x3FF << 16) | 0x02))
    add(t, 2, "pcie", "admin", sqe(0xC1, cid=2, cdw10=0x1234))                        # vendor specific: unknown to the catalog
    # test_mi: MI data structure reads + a tunnelled admin identify
    t = "test_mi"
    add(t, 1, "mi", "", None, mi_msg(nmimt=1, opcode=0x00, nmd0=0x00000000))          # DTYP=00h
    add(t, 2, "mi", "", None, mi_msg(nmimt=1, opcode=0x00, nmd0=0x03000000))          # DTYP=03h
    add(t, 3, "mi", "", None, mi_msg(nmimt=1, opcode=0x01))                           # health status poll
    add(t, 4, "mi", "", None, mi_admin_msg(0x06, nsid=0, cdw10=0x01))                # NVMe admin Identify CNS=01h over MI
    # fill with repeated traffic to exercise signature dedup
    for i in range(200):
        add("test_stress", i, "pcie", "io", sqe(0x02, nsid=1, cid=i & 0xFFFF, cdw10=0, cdw12=7, prp1=0x2000 + i * 0x1000))
    if layout == "A":
        con.executemany("INSERT INTO commands(test_name,script,seq,ts,interface,queue,qid,sqe,mi_msg) VALUES(?,?,?,?,?,?,?,?,?)", rows)
    else:
        for (test, script, seq, ts, iface, queue, qid, sqe_b, mi_b) in rows:
            if sqe_b is None:
                continue
            dws = struct.unpack("<16I", sqe_b)
            con.execute("INSERT INTO nvme_cmds(testcase,script,step,queue,%s) VALUES(?,?,?,?,%s)" %
                        (",".join(f"cdw{i}" for i in range(16)), ",".join("?" * 16)), (test, script, seq, queue, *dws))
    con.commit()
    con.close()
    return truth
